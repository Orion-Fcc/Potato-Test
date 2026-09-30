"""Import an Excel test-case set into a project via the API.

    python scripts/import_xlsx.py testcases/scenarios.xlsx <project_id> [base_url]

Columns expected (header row): name, prompt, expected, start_url, tags, enabled.
`tags` is comma-separated; `enabled` accepts true/false/1/0/yes/no.
"""

from __future__ import annotations

import sys

import httpx
from openpyxl import load_workbook


def _as_bool(v: object) -> bool:
    if isinstance(v, bool):
        return v
    if v is None:
        return True
    return str(v).strip().lower() not in ("false", "0", "no", "")


def load_cases(path: str) -> list[dict]:
    ws = load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    cases: list[dict] = []
    for raw in rows[1:]:
        d = dict(zip(header, raw))
        if not d.get("name"):
            continue
        tags = [t.strip() for t in str(d.get("tags") or "").split(",") if t.strip()]
        cases.append(
            {
                "name": str(d["name"]),
                "prompt": str(d.get("prompt") or ""),
                "expected": str(d.get("expected") or ""),
                "start_url": str(d["start_url"]) if d.get("start_url") else None,
                "tags": tags,
                "enabled": _as_bool(d.get("enabled")),
            }
        )
    return cases


def main(path: str, project_id: int, base: str) -> None:
    cases = load_cases(path)
    r = httpx.post(f"{base}/api/projects/{project_id}/testcases/import", json=cases, timeout=30)
    r.raise_for_status()
    print(f"imported {len(cases)} cases into project {project_id}: {r.json()}")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("usage: python scripts/import_xlsx.py <xlsx> <project_id> [base_url]")
        raise SystemExit(1)
    main(
        sys.argv[1], int(sys.argv[2]), sys.argv[3] if len(sys.argv) > 3 else "http://localhost:8099"
    )
