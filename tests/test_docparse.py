"""文档解析（`app/docparse.py`）的覆盖测试（2026-10-17）。

起因是用户一个很实在的要求：

> 能确保 ai 能阅读到原文档吗，如果不适配，建议改成 md，但是前端我看到的得是原文档，
> 但是 ai 要阅读到所有文档内容

拆成三条约束，这组测试守的就是它们：

1. **AI 要读到全部内容** —— 尤其是 docx 表格（原来只取段落，表格会整个漏掉）；
2. **顺序要对** —— 原来是"先所有段落、再所有表格"，于是"下表列出了…"和表被拆散，
   模型读到的上下文是错位的；
3. **读不全要明说、并给可执行建议（转 md / 看原件）** —— 最危险的不是报错，
   而是"读得很少却一声不响"：界面上传成功，AI 其实什么都没看到。
"""

from __future__ import annotations

import io
import pathlib
import sys

import pytest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

from app import docparse  # noqa: E402


def _docx_bytes() -> bytes:
    """段 → 表 → 段，顺序必须原样保留。"""
    import docx

    d = docx.Document()
    d.add_paragraph("第一段：下面是关键表格。")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "字段"
    t.cell(0, 1).text = "含义"
    t.cell(1, 0).text = "aircraftNo"
    t.cell(1, 1).text = "机号"
    d.add_paragraph("第三段：表格之后的小结。")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_reads_tables_not_only_paragraphs():
    """★ docx 的**表格内容必须读到** —— 需规里一大半信息在表里。"""
    res = docparse.extract("需规.docx", _docx_bytes())
    assert "aircraftNo" in res.text and "机号" in res.text
    assert res.fmt == "docx"


def test_docx_keeps_document_order():
    """★ 段落与表格按**原文顺序**出现。

    原来先输出所有段落、再输出所有表格，于是"下面是关键表格"后面跟着的是第三段 ——
    AI 读到的上下文是错的。这里直接断言三者出现的先后。
    """
    res = docparse.extract("需规.docx", _docx_bytes())
    i1 = res.text.index("第一段")
    it = res.text.index("aircraftNo")
    i3 = res.text.index("第三段")
    assert i1 < it < i3, res.text[:400]


def test_scanned_pdf_page_is_reported_with_advice():
    """没有文字层的页要**点名**并给建议 —— 这才是"AI 读不到"的真实场景。"""
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    w.write(buf)

    warns: list[str] = []
    docparse._from_pdf(buf.getvalue(), warns)  # 直接测抽取器：全空页会被上层当作"无文字"
    assert warns, "扫描页一声不响 —— 用户会以为 AI 读到了"
    joined = " ".join(warns)
    assert "没有文字层" in joined
    assert "md" in joined, "只说了读不到，没给可执行的建议"


def test_big_file_with_little_text_is_flagged():
    """★ 覆盖率兜底：文件不小、却只读到几十个字 —— 必须报警（多半是图片/扫描件）。

    这条是"确保 AI 读到所有内容"的最后一道网：不报错但内容缺失最容易被忽略。
    """
    data = b"abc" + b" " * 60_000  # 约 58 KB，只有 3 个可见字符
    res = docparse.extract("扫描件.txt", data)
    assert any("只读到" in w and "md" in w for w in res.warnings), res.warnings


def test_small_file_does_not_cry_wolf():
    """正常的小文件**不能**报这个警 —— 否则用户会习惯性忽略所有提示。"""
    res = docparse.extract("说明.md", "# 标题\n\n很短但正常的一段。\n".encode("utf-8"))
    assert not res.warnings, res.warnings


def test_unsupported_binary_points_at_md():
    """完全不认识的二进制要告诉用户"先转 md"，而不是只丢一句 unsupported。"""
    with pytest.raises(docparse.UnsupportedDocument) as ei:
        docparse.extract("图片.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 512)
    assert "md" in str(ei.value).lower()


def test_warnings_are_written_for_humans():
    """告警是给**用户**看的（会以 toast 弹出），必须是中文、带建议。

    原来这些文字是英文（"the .docx has no extractable text…"），中文界面上蹦出英文
    等于没提示 —— 这条防回归。
    """
    import docx

    d = docx.Document()  # 一个没有任何内容的 docx
    buf = io.BytesIO()
    d.save(buf)

    warns: list[str] = []
    docparse._from_docx(buf.getvalue(), warns)
    assert warns
    assert "没读到任何文字" in warns[0]
    assert "md" in warns[0]
